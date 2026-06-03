import csv
import io
import shutil
from datetime import date, datetime
from pathlib import Path

from flask import current_app
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .db import get_db
from .formula_engine import calculate_materials


def today_iso():
    return date.today().isoformat()


def upsert_named(table, name, **extra):
    if not name:
        return
    db = get_db()
    columns = ["name", *extra.keys()]
    values = [name, *extra.values()]
    placeholders = ",".join(["?"] * len(values))
    updates = ",".join([f"{key}=excluded.{key}" for key in extra])
    sql = f"INSERT INTO {table} ({','.join(columns)}) VALUES ({placeholders}) ON CONFLICT(name) DO "
    sql += f"UPDATE SET {updates}" if updates else "NOTHING"
    db.execute(sql, values)


def list_projects(filters=None):
    db = get_db()
    filters = filters or {}
    where = []
    params = []
    if filters.get("state"):
        state = filters["state"]
        if state == "completed":
            where.append("p.work_done = 1 AND (p.final_price - COALESCE(pay.total_paid, 0)) <= 0")
        elif state == "unpaid":
            where.append("p.work_done = 1 AND (p.final_price - COALESCE(pay.total_paid, 0)) > 0")
        elif state == "in_progress":
            where.append("p.work_done = 0")
        elif state == "overdue":
            where.append("p.work_done = 0 AND p.deadline IS NOT NULL AND p.deadline < ?")
            params.append(today_iso())
    if filters.get("contractor"):
        where.append("p.contractor = ?")
        params.append(filters["contractor"])
    if filters.get("from"):
        where.append("p.date >= ?")
        params.append(filters["from"])
    if filters.get("to"):
        where.append("p.date <= ?")
        params.append(filters["to"])
    search = filters.get("search")
    if search:
        where.append("(p.client_name LIKE ? OR p.shop_name LIKE ? OR p.phone LIKE ? OR p.project_type LIKE ? OR p.contractor LIKE ?)")
        params.extend([f"%{search}%"] * 5)
    sql = """
        SELECT p.*,
               COALESCE(mat.materials_total, 0) AS materials_total,
               COALESCE(mat.labor_total, 0) AS labor_total,
               COALESCE(pay.total_paid, 0) AS total_paid,
               ROUND(p.final_price - COALESCE(pay.total_paid, 0), 2) AS remaining_balance,
               ROUND(p.final_price - COALESCE(mat.materials_total, 0) - COALESCE(mat.labor_total, 0), 2) AS profit
        FROM projects p
        LEFT JOIN (
            SELECT project_id,
                   SUM(CASE WHEN is_labor = 1 THEN 0 ELSE total_price END) materials_total,
                   SUM(CASE WHEN is_labor = 1 THEN total_price ELSE 0 END) labor_total
            FROM project_materials GROUP BY project_id
        ) mat ON mat.project_id = p.id
        LEFT JOIN (SELECT project_id, SUM(amount) total_paid FROM payments GROUP BY project_id) pay ON pay.project_id = p.id
    """
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY p.date DESC, p.id DESC"
    return db.execute(sql, params).fetchall()


def get_project(project_id):
    db = get_db()
    project = db.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
    if not project:
        return None
    materials = db.execute("SELECT * FROM project_materials WHERE project_id = ? ORDER BY id", (project_id,)).fetchall()
    payments = db.execute("SELECT * FROM payments WHERE project_id = ? ORDER BY payment_date DESC, id DESC", (project_id,)).fetchall()
    images = db.execute("SELECT * FROM project_images WHERE project_id = ? ORDER BY is_primary DESC, id", (project_id,)).fetchall()
    totals = project_totals(project_id)
    return {**project, "materials": materials, "payments": payments, "images": images, "totals": totals}


def save_project(payload, user_id=None, project_id=None):
    db = get_db()
    fields = {
        "project_type": payload.get("project_type", "").strip(),
        "date": payload.get("date") or today_iso(),
        "deadline": payload.get("deadline"),
        "shop_name": payload.get("shop_name", "").strip(),
        "client_name": payload.get("client_name", "").strip(),
        "phone": payload.get("phone", "").strip(),
        "address": payload.get("address", "").strip(),
        "contractor": payload.get("contractor", "").strip(),
        "width": float(payload.get("width") or 0),
        "height": float(payload.get("height") or 0),
        "notes": payload.get("notes", "").strip(),
        "work_done": 1 if payload.get("work_done") else 0,
        "status": "completed" if payload.get("work_done") else "in progress",
        "final_price": float(payload.get("final_price") or 0),
    }
    if not fields["client_name"]:
        raise ValueError("Client name is required")

    for table, value in [("clients", fields["client_name"]), ("shops", fields["shop_name"]), ("contractors", fields["contractor"]), ("project_types", fields["project_type"])]:
        upsert_named(table, value, phone=fields["phone"] if table == "clients" else None) if table == "clients" else upsert_named(table, value)

    if project_id:
        assignments = ", ".join([f"{key}=?" for key in fields])
        db.execute(f"UPDATE projects SET {assignments}, updated_at=CURRENT_TIMESTAMP WHERE id=?", [*fields.values(), project_id])
        db.execute("DELETE FROM project_materials WHERE project_id=?", (project_id,))
    else:
        columns = [*fields.keys(), "created_by"]
        placeholders = ",".join(["?"] * len(columns))
        cursor = db.execute(f"INSERT INTO projects ({','.join(columns)}) VALUES ({placeholders})", [*fields.values(), user_id])
        project_id = cursor.lastrowid

    materials = calculate_materials(fields["width"], fields["height"], payload.get("materials", []))
    for material in materials:
        db.execute(
            """
            INSERT INTO project_materials
            (project_id, material_id, material_name, category, unit, quantity, unit_price, formula, auto_formula, manual_override, waste_percent, is_labor, total_price)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                project_id,
                material.get("material_id"),
                material.get("material_name", "").strip(),
                material.get("category", ""),
                material.get("unit", "Pcs"),
                material.get("quantity") or 0,
                material.get("unit_price") or 0,
                material.get("formula", ""),
                1 if material.get("auto_formula") else 0,
                1 if material.get("manual_override") else 0,
                material.get("waste_percent") or 0,
                1 if material.get("is_labor") or material.get("category") == "Labor" or material.get("material_name") == "Working hand" else 0,
                material.get("total_price") or 0,
            ),
        )
        remember_template(fields["project_type"], material)
    db.commit()
    return get_project(project_id)


def remember_template(project_type, material):
    if not project_type or not material.get("material_name"):
        return
    get_db().execute(
        """
        INSERT INTO material_templates
        (project_type, material_id, material_name, category, unit, unit_price, formula, auto_formula, waste_percent)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(project_type, material_name) DO UPDATE SET
        material_id=excluded.material_id, category=excluded.category, unit=excluded.unit,
        unit_price=excluded.unit_price, formula=excluded.formula,
        auto_formula=excluded.auto_formula, waste_percent=excluded.waste_percent
        """,
        (
            project_type,
            material.get("material_id"),
            material.get("material_name"),
            material.get("category", ""),
            material.get("unit", "Pcs"),
            material.get("unit_price") or 0,
            material.get("formula", ""),
            1 if material.get("auto_formula") else 0,
            material.get("waste_percent") or 0,
        ),
    )


def project_totals(project_id):
    db = get_db()
    project = db.execute("SELECT final_price FROM projects WHERE id = ?", (project_id,)).fetchone()
    if not project:
        return {}
    materials_total = db.execute("SELECT COALESCE(SUM(total_price), 0) AS value FROM project_materials WHERE project_id = ? AND is_labor = 0", (project_id,)).fetchone()["value"]
    labor_total = db.execute("SELECT COALESCE(SUM(total_price), 0) AS value FROM project_materials WHERE project_id = ? AND is_labor = 1", (project_id,)).fetchone()["value"]
    paid = db.execute("SELECT COALESCE(SUM(amount), 0) AS value FROM payments WHERE project_id = ?", (project_id,)).fetchone()["value"]
    materials_total = round(materials_total or 0, 2)
    final_price = round(project["final_price"] or 0, 2)
    paid = round(paid or 0, 2)
    return {
        "materials_total": materials_total,
        "price": round(materials_total * 2, 2),
        "final_price": final_price,
        "labor_total": round(labor_total or 0, 2),
        "profit": round(final_price - materials_total - (labor_total or 0), 2),
        "paid": paid,
        "remaining_balance": round(final_price - paid, 2),
    }

def stats(filters=None):
    projects = list_projects(filters)
    total_profit = sum(p["profit"] for p in projects)
    unpaid = sum(max(p["remaining_balance"], 0) for p in projects if p["status"] != "paid")
    completed = sum(1 for p in projects if p["status"] in ("completed", "paid"))
    pending = sum(1 for p in projects if p["status"] == "in progress")
    months = {p["date"][:7] for p in projects if p.get("date")}
    return {
        "total_transactions": len(projects),
        "total_profit": round(total_profit, 2),
        "unpaid_amount": round(unpaid, 2),
        "average_monthly_income": round(sum(p["final_price"] for p in projects) / max(len(months), 1), 2),
        "completed_jobs": completed,
        "pending_jobs": pending,
        "by_month": monthly_chart(projects),
        "status_mix": status_chart(projects),
    }


def monthly_chart(projects):
    data = {}
    for project in projects:
        key = project["date"][:7]
        data[key] = data.get(key, 0) + project["final_price"]
    return [{"label": key, "value": round(value, 2)} for key, value in sorted(data.items())]


def status_chart(projects):
    data = {}
    for project in projects:
        data[project["status"]] = data.get(project["status"], 0) + 1
    return [{"label": key, "value": value} for key, value in data.items()]


def create_backup():
    src = Path(current_app.config["DATABASE"])
    dst_dir = Path(current_app.config["BACKUP_DIR"])
    dst_dir.mkdir(parents=True, exist_ok=True)
    dst = dst_dir / f"workshop-{datetime.now().strftime('%Y%m%d-%H%M%S')}.sqlite3"
    shutil.copy2(src, dst)
    return dst


def export_projects(fmt):
    projects = list_projects({})
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=projects[0].keys() if projects else ["id"])
        writer.writeheader()
        writer.writerows(projects)
        return buffer.getvalue().encode("utf-8"), "text/csv", "projects.csv"
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Projects"
        headers = list(projects[0].keys()) if projects else ["id"]
        ws.append(headers)
        for project in projects:
            ws.append([project.get(h) for h in headers])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "projects.xlsx"
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(40, y, "Workshop Projects")
    pdf.setFont("Helvetica", 9)
    for project in projects[:45]:
        y -= 16
        pdf.drawString(40, y, f"{project['date']} | {project['client_name']} | {project['project_type'] or ''} | {project['final_price']:.2f}")
    pdf.save()
    return buffer.getvalue(), "application/pdf", "projects.pdf"
